import docx
import os

# Example script to generate Fiches de description
def create_fiche(module_id, module_title, eval_type, conditions, description, criteres, output_dir):
    doc = docx.Document("Fiche de description d.docx")
    table = doc.tables[0]
    
    # Common fields
    table.cell(0, 0).text = module_id
    table.cell(2, 2).text = module_title
    table.cell(3, 0).text = "Module " + module_id.split('-')[-1]
    table.cell(3, 2).text = f"Objectifs : Acquérir les compétences liées au module {module_title}."
    table.cell(5, 0).text = "Date :"
    table.cell(5, 1).text = "2 Heures"
    table.cell(5, 2).text = eval_type
    
    # Specific fields
    table.cell(7, 0).text = "Conditions et organisation du contrôle :\n" + conditions
    table.cell(7, 4).text = "Critères d’appréciation des résultats :\n" + criteres
    table.cell(8, 0).text = "Description des objectifs de l’épreuve :\n" + description
    
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
        
    filename = f"Fiche {module_id} - {eval_type}.docx"
    doc.save(os.path.join(output_dir, filename))
    print(f"Created: {filename}")

# Usage Example:
# create_fiche("ATAM-01", "Métier et formation", "Contrôle Continu", "...", "...", "...", "./output")
